from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "submission" / "SME-Operating-Agent-Business-Proposal-Draft.docx"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
PALE_GRAY = "F6F7F8"
GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), LIGHT_GRAY)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Aptos", size=None, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_body(doc, text="", bold_lead=None, italic=False, after=7, keep=False):
    p = doc.add_paragraph()
    p.style = doc.styles["Body Text"]
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, italic=italic)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = keep
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.27 + 0.22 * level)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        bullet = p.add_run("•  ")
        set_run_font(bullet)
        r = p.add_run(item)
        set_run_font(r)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True


def add_numbered(doc, items):
    for number, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.31)
        p.paragraph_format.first_line_indent = Inches(-0.27)
        marker = p.add_run(f"{number}.  ")
        set_run_font(marker)
        r = p.add_run(item)
        set_run_font(r)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True


def add_table(doc, headers, rows, widths=None, font_size=9.4, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_run_font(r, size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        fill = WHITE if row_index % 2 == 0 else PALE_BLUE
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    if getattr(doc, "_pending_page_break", False):
        p.paragraph_format.page_break_before = True
        doc._pending_page_break = False
    return p


def page_break(doc):
    doc._pending_page_break = True


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    body = styles["Body Text"]
    body.font.name = "Aptos"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    body.font.size = Pt(10.8)
    body.font.color.rgb = RGBColor.from_string(BLACK)
    body.paragraph_format.line_spacing = 1.12
    body.paragraph_format.space_after = Pt(7)

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(31)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 19, 12, 8), ("Heading 2", 13.5, 9, 5), ("Heading 3", 11.5, 7, 4)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.5)

    return doc


def build():
    doc = setup_document()

    # Cover
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(74)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("SME Operating Agent Business Proposal")
    set_run_font(r, name="Aptos Display", size=31, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("A practical agent for company knowledge meeting follow up and founder led recruiting")
    set_run_font(r, name="Aptos Display", size=17, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Business Proposal Draft")
    set_run_font(r, size=12, bold=True, color=NAVY)

    cover_table = doc.add_table(rows=5, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cover_table.autofit = False
    labels = ["Company", "Team code", "Submission", "Repository", "Date"]
    values = [
        "Stellar Ark AI",
        "Pending confirmation",
        "Show Me Your Agents Hackathon",
        "github.com/joshualumzy/personal_agent_long_context",
        "27 September 2026",
    ]
    for idx, (label, value) in enumerate(zip(labels, values)):
        a, b = cover_table.rows[idx].cells
        a.width = Inches(1.35)
        b.width = Inches(5.15)
        for cell in (a, b):
            set_cell_margins(cell, top=65, start=0, bottom=65, end=80)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        a.text = ""
        b.text = ""
        ar = a.paragraphs[0].add_run(label)
        set_run_font(ar, size=9.5, bold=True, color=GRAY)
        br = b.paragraphs[0].add_run(value)
        set_run_font(br, size=10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Proposal purpose")
    set_run_font(r, size=10, bold=True, color=NAVY)
    add_body(
        doc,
        "Secure approval to pilot one integrated agent that helps a small team recover company context, convert meeting commitments into controlled follow up, and run a structured recruiting workflow.",
        after=0,
    )

    page_break(doc)

    heading(doc, "Executive Summary")
    add_body(
        doc,
        "Stellar Ark AI builds AI-native wearables, including context glasses and a pendant, paired with a proactive agent system that turns everyday conversations into persistent memory and action. Its agent layer connects with Feishu, email, calendar, and CRM to generate meeting minutes, update customer profiles, and synchronise action items for professionals in consulting, finance, legal, enterprise sales, and technology.",
    )
    add_body(
        doc,
        "The company faces a related internal operating challenge. Employees must recover company context from scattered records, convert meeting decisions into accountable follow up, and support hiring while continuing core product and customer work. The information needed for those tasks spans messages, documents, meetings, business systems, and individual memory.",
    )
    add_body(
        doc,
        "We propose the SME Operating Agent, a single product with three connected workflows. Company Context retrieves employee-visible evidence and returns inspectable citations. Meeting Actions converts text transcript lines into traceable answers, conflicts, and reviewable drafts. Recruiting helps a founder define criteria, search and assess public professional profiles, and prepare outreach without sending on the founder's behalf.",
    )
    add_body(
        doc,
        "The current prototype demonstrates the product architecture and human-control model using privacy-safe synthetic data. It does not expose Stellar Ark AI business records and does not yet establish business savings or production reliability. The recommended next step is a bounded internal pilot that measures time saved, follow-through, quality, and control failures against the team's current workflow.",
    )

    heading(doc, "One Business Story", 2)
    add_table(
        doc,
        ["Know", "Act", "Grow"],
        [[
            "Recover company context from inspectable evidence",
            "Turn text meetings into controlled follow up work",
            "Run a structured founder led recruiting workflow",
        ]],
        widths=[2.15, 2.15, 2.15],
        font_size=10,
    )
    add_body(
        doc,
        "These workflows share one operating principle: the agent may retrieve, analyse, and draft, but important external actions remain visible to a person. Company questions use cited evidence. Meeting drafts require the correct policy tier and exact-payload approval. Recruiting outreach stays in the founder's own Gmail or LinkedIn account.",
    )

    heading(doc, "Recommendation", 2)
    add_body(
        doc,
        "Proceed to a four to six week internal pilot at Stellar Ark AI with one team, one recurring meeting type, and one active hiring role. Establish baselines before activation, configure only the integrations needed for the pilot, and treat the proposed KPI targets in this document as acceptance hypotheses rather than achieved results.",
    )

    page_break(doc)

    heading(doc, "Problem and Opportunity")
    heading(doc, "Company Context", 2)
    add_body(
        doc,
        "Stellar Ark AI is an SME building AI-native wearables and proactive agent software. Its products turn captured conversations into persistent memory and action through a layered memory architecture and connections to Feishu, email, calendar, and CRM. The hackathon team is participating from within this business context, so the proposed problem is grounded in first-hand operational understanding rather than an invented target company.",
    )
    heading(doc, "Target Users", 2)
    add_body(
        doc,
        "The primary users are Stellar Ark AI employees and founders who must turn scattered company information into daily work while contributing to a fast-moving AI-native product business. Managers, IT administrators, and governance owners are secondary stakeholders because they need reliable sources, clear permissions, and visible accountability. The same workflow may later apply to comparable SMEs with limited operations or recruiting capacity.",
    )

    heading(doc, "Current Workflow", 2)
    add_numbered(doc, [
        "An employee receives a company question, makes a commitment in a meeting, or a founder decides to hire.",
        "They search several systems and reconstruct the relevant history by hand.",
        "They translate what they found into an answer, email, ticket, calendar event, document, or hiring criteria.",
        "They chase missing details and remember which action still needs follow up.",
        "They personally check and execute the consequential action.",
    ])

    heading(doc, "Why the Problem Matters", 2)
    add_table(
        doc,
        ["Operational problem", "Business consequence", "Evidence status"],
        [
            ["Fragmented company records", "Slower answers and decisions based on incomplete context", "First-hand company context; quantitative baseline required"],
            ["Meeting commitments rely on memory", "Follow up is delayed, duplicated, or missed", "First-hand workflow basis; frequency and loss rate required"],
            ["Founder led hiring has little support", "Search, screening, and follow up compete with core leadership work", "First-hand workflow basis; hiring-time baseline required"],
        ],
        widths=[2.0, 2.65, 1.8],
        font_size=9.2,
    )
    add_body(
        doc,
        "The opportunity is to reduce the manual gap between finding context and preparing the next piece of work. The proposal does not assume that every task should be automated. It focuses on the repetitive reconstruction, drafting, and tracking steps while retaining human judgment for decisions and external effects.",
    )

    heading(doc, "Problem Basis and Data Approach", 2)
    add_body(
        doc,
        "The problem is based on the team's first-hand understanding of Stellar Ark AI. Private operational records, customer information, and conversations are intentionally excluded from the hackathon. OrgForge provides privacy-safe synthetic demo and evaluation data; it is not a customer, dependency, or integration. A controlled internal pilot will measure handling time, manual steps, missed follow ups, rework, and hiring effort, while documenting permissions, retention, and approval practices before private data is introduced.",
    )

    heading(doc, "Proposed Solution")
    heading(doc, "Company Context", 2)
    add_body(
        doc,
        "An authenticated employee asks a company question. The agent searches employee-visible Company Evidence using full-text retrieval, vector similarity, and explicit links between related records. It returns a concise answer with inspectable source identifiers or states that the available evidence is insufficient. Employee-specific Personal Memory is displayed separately and is not treated as company evidence.",
    )

    heading(doc, "Meeting Actions", 2)
    add_body(
        doc,
        "The meeting workflow accepts transcript text entered live, pasted, or replayed from scripted demo scenarios. It does not perform speech-to-text transcription. Each segment is screened before model use. The service identifies commitments, questions, and decisions, keeps the exact trigger quote, answers company questions through Company Context, checks for conflicting decisions, and prepares supported drafts. OrgForge supplies synthetic hackathon demo and evaluation data only; it is not a customer, product dependency, deployment target, or external integration.",
    )
    add_table(
        doc,
        ["Policy tier", "Examples", "Control"],
        [
            ["Automatic", "Read-only answer or conflict flag", "Executed and traced without changing an external system"],
            ["Approval", "Email, message, calendar, ticket, document, spreadsheet, hiring request", "Approval is bound to the exact payload; an edit requires approval again"],
            ["Escalation", "Money or contract commitment", "The employee cannot execute it; the action is raised to a named authority"],
            ["Blocked", "Secret or transcript instruction attack", "Stopped before model use"],
        ],
        widths=[1.1, 2.4, 3.0],
        font_size=9.2,
    )

    heading(doc, "Recruiting", 2)
    add_body(
        doc,
        "A founder states a hiring need or uploads a job description. The agent proposes reviewable criteria, searches public professional profiles through Exa when configured or uses clearly labelled fictional sample profiles, and scores candidates against the confirmed criteria. Founder feedback can change criteria or create a proposal, but the change applies only after an explicit decision. The agent may find a work email through configured lookup services and draft outreach; it never guesses an address and cannot send through its chat tools.",
    )

    heading(doc, "How the Workflows Connect", 2)
    add_bullets(doc, [
        "Company Context supplies evidence and answers to Meeting Actions.",
        "An approved hiring need from a meeting creates a draft role in Recruiting.",
        "Recruiting is also available from the main company chat as a skill that loads only for hiring requests.",
        "Personal Memory retains employee context and stated hiring rationale through separate adapters, while Company Evidence remains in PostgreSQL.",
    ])

    page_break(doc)

    heading(doc, "Business Value")
    add_body(
        doc,
        "The business case depends on measurable reductions in manual work and avoidable rework. The prototype supports the workflow, but the following benefits remain hypotheses until an internal pilot measures them against Stellar Ark AI's baseline.",
    )
    add_table(
        doc,
        ["Value area", "How value is created", "Pilot evidence required"],
        [
            ["Productivity", "Fewer manual searches and less time converting context into drafts", "Median handling time and manual steps before and after"],
            ["Follow through", "Meeting commitments become visible, traceable work instead of relying on memory", "Completion rate, time to first draft, and overdue actions"],
            ["Hiring capacity", "A founder receives structured criteria, a reviewable pool, and follow up drafts", "Founder hours per role and time from requirement to reviewed pool"],
            ["Service consistency", "Answers use inspectable evidence and unsupported questions can stop explicitly", "Correctness, evidence recall, and user source inspection"],
            ["Risk control", "Deterministic policy, exact-payload approval, and user-account handoffs limit silent action", "Approval bypasses, stale approvals, unsafe content, and outbound-send violations"],
            ["Scale", "The same evidence and human-control foundation supports several workflows", "Users, cases, latency, operating cost, and support effort"],
        ],
        widths=[1.15, 3.25, 2.1],
        font_size=8.9,
    )

    heading(doc, "Before and After", 2)
    add_table(
        doc,
        ["Task", "Current process", "Proposed process"],
        [
            ["Company question", "Search systems, connect records, write an answer", "Retrieve evidence, inspect citations, accept the answer or insufficient-evidence result"],
            ["Meeting commitment", "Take notes, interpret the promise, gather details, prepare work later", "Screen transcript text, identify the commitment, draft now, review the exact payload, hand off"],
            ["Founder hiring", "Write criteria, search profiles, compare candidates, remember feedback, chase replies", "Confirm criteria, review a scored pool, record feedback, prepare outreach and follow ups"],
        ],
        widths=[1.35, 2.55, 2.6],
        font_size=9.1,
    )

    heading(doc, "Commercial Position", 2)
    add_body(
        doc,
        "No revenue, cost-saving, or return-on-investment claim is established today. A pilot should first determine whether the saved staff and founder time exceeds the operating cost of model calls, storage, integrations, onboarding, and support. Pricing and packaging should follow that evidence rather than precede it.",
    )

    page_break(doc)

    heading(doc, "Impact and Outcomes")
    add_body(
        doc,
        "The pilot should compare representative cases before and after the agent is introduced. Proposed targets below are decision thresholds for the pilot, not results already achieved. Baselines must be measured during the first week and targets adjusted if the task mix differs materially from the assumptions.",
    )
    add_table(
        doc,
        ["Outcome", "Metric", "Proposed pilot target"],
        [
            ["Faster context reconstruction", "Median time to produce an evidence-backed answer", "At least 30 percent lower than baseline"],
            ["Less meeting administration", "Median time from meeting end to reviewable follow up draft", "At least 50 percent lower than baseline"],
            ["Better follow through", "Share of accepted meeting actions completed by their due date", "At least 90 percent"],
            ["Lower founder hiring effort", "Founder administration time per active role", "At least 25 percent lower than baseline"],
            ["Trustworthy company answers", "Factual accuracy and expected-evidence recall on pilot questions", "Threshold set after baseline; both must improve from the current prototype"],
            ["Human control", "Unauthorized outbound sends or executed escalated actions", "Zero"],
            ["Usable evidence", "Pilot users who can open and understand the source behind an answer", "At least 90 percent"],
        ],
        widths=[1.7, 3.0, 1.8],
        font_size=8.9,
    )

    heading(doc, "Measurement Plan", 2)
    add_numbered(doc, [
        "Capture baseline duration, steps, rework, and completion for a representative set of company questions, meetings, and one hiring role.",
        "Run the same categories through the agent and log retrieval, model, policy, approval, handoff, error, and completion timestamps.",
        "Review answer correctness and source support with an authorised SME owner; review hiring judgments for relevance and fairness with a human hiring owner.",
        "Conduct weekly user interviews and inspect failure cases, not only aggregate scores.",
        "At the end of the pilot, compare value, risk, operating cost, and support effort before deciding whether to expand.",
    ])

    page_break(doc)
    heading(doc, "Current Technical Evidence", 2)
    add_table(
        doc,
        ["Area", "Recorded result", "Interpretation"],
        [
            ["Company Context", "76-question run: 12 percent factual accuracy, 29 percent expected-evidence recall, 100 percent citation-ID integrity, 39.68 seconds mean latency", "Citation identifiers were controlled in that sample, but answer quality and latency require improvement"],
            ["Meeting Actions", "One live run: 23 of 23 policy tiers correct; 9 of 9 unsafe lines blocked; 0 of 44 ordinary lines blocked; 1 of 1 conflict found; email recall 2 of 3", "Encouraging bounded evidence from one model run, not a reliability guarantee"],
            ["Recruiting", "Deterministic, holdout, regression, chaos, and browser test assets exist", "A final submission-snapshot quality summary remains to be produced"],
        ],
        widths=[1.2, 3.3, 2.0],
        font_size=8.5,
    )
    add_body(
        doc,
        "These results support a pilot recommendation, not a production claim. The company-answer evaluation is the clearest current quality gap. The final submission should show improvement or explain why the pilot is designed to address it.",
    )

    page_break(doc)

    heading(doc, "Feasibility and Scalability")
    heading(doc, "What Exists Today", 2)
    add_bullets(doc, [
        "A TypeScript and Fastify modular monolith with three registered browser surfaces.",
        "PostgreSQL and pgvector for Company Evidence, conversations, employee records, meeting state, and action logs.",
        "A multi-step company reasoning loop with source validation and Server-Sent Events.",
        "A meeting service with pre-model screening, extraction, evidence gathering, deterministic policy, payload versioning, approval, and handoffs.",
        "A per-role recruiting service with criteria, search, scoring, feedback, proposals, outreach drafts, reply handling, retention, and a main-chat skill.",
        "Letta adapters for employee context and founder hiring-intent memory, separate from Company Evidence.",
    ])

    heading(doc, "Operating Requirements", 2)
    add_table(
        doc,
        ["Capability", "Required for core pilot", "Optional live integration"],
        [
            ["Application", "Node.js runtime, Fastify server, session secret", "Public deployment and managed secrets"],
            ["Company evidence", "PostgreSQL with pgvector and permission-scoped data", "Incremental connectors for approved systems"],
            ["Models and memory", "SoCLaaS or supported model path; Letta or documented fallback", "Provider redundancy and monitored service levels"],
            ["Recruiting", "Labelled sample profiles or Exa; manual outreach completion", "Exa, Hunter, Prospeo, Google OAuth, read-only LinkedIn sync"],
            ["Meeting handoffs", "Browser links and explicit simulated state where required", "Configured GitHub repository and approved office-suite choices"],
        ],
        widths=[1.35, 2.7, 2.45],
        font_size=9.0,
    )

    heading(doc, "Path to Scale", 2)
    add_numbered(doc, [
        "Prototype to pilot: select one SME, one department, one meeting type, and one active hiring role.",
        "Pilot hardening: add permission rules, consent and retention controls, monitoring, support procedures, and a bounded evaluation set.",
        "Production transition: replace demo personas and local recruiting files with tenant-scoped identity, authorisation, and durable storage.",
        "Controlled expansion: add connector-specific ingestion and approval contracts only after each workflow meets its quality and safety threshold.",
        "Cross-functional scale: reuse the shared evidence and human-control foundation for other departments while preserving separate permissions and measures.",
    ])

    heading(doc, "Scalability Questions to Resolve", 2)
    add_bullets(doc, [
        "Expected user count, case volume, latency requirement, and model budget",
        "Permission model across departments and data sources",
        "Operational ownership for monitoring, incident response, backup, and deletion",
        "Legal basis and fairness controls for recruiting data and decisions",
    ])

    heading(doc, "Risk Governance and Human Control")
    add_table(
        doc,
        ["Risk", "Current control", "Pilot requirement"],
        [
            ["Incorrect answer with valid-looking sources", "Inspectable evidence and citation-ID validation", "Human-labelled answer and semantic source-support review"],
            ["Missing evidence", "Full-text, vector, and explicit-link retrieval; insufficient-evidence outcome", "Improve recall against pilot questions before expansion"],
            ["Unsafe meeting content", "Pre-model screening for prohibited data and instruction attacks", "Adversarial testing with SME-specific examples"],
            ["Stale or unauthorised action", "Deterministic tier and exact-payload approval; edits invalidate approval", "Server-derived identity, audit review, and negative approval tests"],
            ["Silent external send", "S2 and S3 prepare handoffs; S3 chat tools expose no send operation", "Confirm zero outbound sends without a user's action"],
            ["Recruiting privacy or bias", "Public professional fields, no guessed email, private pass reasons blocked, closed-candidate retention rule", "Legal review, fairness measures, consent basis, and deletion verification"],
            ["Evaluation leakage", "Reserved oracle records excluded from the runtime Company Evidence path", "Re-verify the final deployed dataset and access boundary"],
        ],
        widths=[1.55, 2.7, 2.25],
        font_size=8.6,
    )

    heading(doc, "Claim Boundary", 2)
    add_body(
        doc,
        "The proposal may claim that the current code integrates three workflows, implements the stated controls, and addresses a problem grounded in the team's first-hand Stellar Ark AI context. It may report saved evaluation results with their sample and run limitations. It should not claim quantitative validation on private company data, production readiness, complete prompt-injection resistance, unbiased hiring, autonomous transcription, autonomous sending, or proven financial returns.",
    )

    heading(doc, "Pilot Decision", 2)
    add_body(
        doc,
        "Approve a bounded pilot if the SME can provide an authorised evidence set, a recurring meeting workflow, one active hiring role, and owners for security and outcome review. Do not expand beyond the pilot until answer quality improves, human-control tests pass, recruiting governance is agreed, and operating cost is measured.",
    )

    page_break(doc)
    heading(doc, "Information Required Before Submission", 2)
    add_table(
        doc,
        ["Item", "Why it matters"],
        [
            ["Team code and final product name", "Required submission identity"],
            ["Quantified Stellar Ark AI workflow baseline", "Supports the opportunity and measurable-outcome claims"],
            ["Final deployment URL and evidence", "Shows that the prototype can be accessed and operated"],
            ["Final verification and S3 evaluation summary", "Aligns technical evidence with the proposal and demo"],
            ["Demo scenario and configured integrations", "Ensures the proposal and recorded demonstration tell the same story"],
        ],
        widths=[2.5, 4.0],
        font_size=9.2,
    )

    heading(doc, "Conclusion", 2)
    add_body(
        doc,
        "The SME Operating Agent addresses a connected operating problem at Stellar Ark AI: employees have to recover context, convert decisions into work, and sustain hiring while continuing core product and customer work. The current implementation is credible enough for a controlled internal pilot because it joins the three workflows and keeps consequential actions visible to people. The business case now depends on quantifying the baseline, improving company-answer quality, and measuring whether the saved effort justifies the operating cost and governance burden.",
        after=0,
    )

    # Footer with a simple proposal label and page number field.
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        r = p.add_run("SME Operating Agent Business Proposal   |   ")
        set_run_font(r, size=8, color=GRAY)
        fld_char1 = OxmlElement("w:fldChar")
        fld_char1.set(qn("w:fldCharType"), "begin")
        instr_text = OxmlElement("w:instrText")
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = " PAGE "
        fld_char2 = OxmlElement("w:fldChar")
        fld_char2.set(qn("w:fldCharType"), "end")
        field_run = p.add_run()
        field_run._r.append(fld_char1)
        field_run._r.append(instr_text)
        field_run._r.append(fld_char2)
        set_run_font(field_run, size=8, color=GRAY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "SME Operating Agent Business Proposal"
    doc.core_properties.subject = "Show Me Your Agents Hackathon business proposal draft"
    doc.core_properties.author = "SME Operating Agent Team"
    doc.core_properties.keywords = "SME, agent, business proposal, company context, meeting actions, recruiting"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
