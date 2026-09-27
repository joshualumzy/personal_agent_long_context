from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_business_proposal import (
    BLACK,
    BLUE,
    GRAY,
    NAVY,
    PALE_BLUE,
    WHITE,
    add_body,
    add_bullets,
    add_numbered,
    add_table,
    heading,
    page_break,
    set_cell_margins,
    set_run_font,
    setup_document,
)


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "submission" / "SME-Operating-Agent-Technical-Documentation-Draft.docx"
DIAGRAM = ROOT / "tmp" / "proposal" / "technical-architecture.png"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/SFNS.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline, title, lines, title_size=30, body_size=23):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=3)
    draw.text((x1 + 24, y1 + 19), title, font=font(title_size, True), fill="#000000")
    y = y1 + 66
    for line in lines:
        draw.text((x1 + 24, y), line, font=font(body_size), fill="#222222")
        y += body_size + 10


def arrow(draw, start, end):
    draw.line([start, end], fill="#17365D", width=7)
    x2, y2 = end
    if start[1] == end[1]:
        direction = 1 if x2 > start[0] else -1
        points = [(x2, y2), (x2 - 18 * direction, y2 - 13), (x2 - 18 * direction, y2 + 13)]
    else:
        direction = 1 if y2 > start[1] else -1
        points = [(x2, y2), (x2 - 13, y2 - 18 * direction), (x2 + 13, y2 - 18 * direction)]
    draw.polygon(points, fill="#17365D")


def build_architecture_diagram():
    DIAGRAM.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1500, 930), "white")
    draw = ImageDraw.Draw(image)

    draw.text((40, 24), "SME Operating Agent Runtime Architecture", font=font(38, True), fill="#000000")
    draw.text((40, 72), "One modular monolith with three browser workflows and separated evidence, memory, and action state", font=font(23), fill="#555555")

    rounded_box(draw, (55, 140, 465, 315), "#DCE6F1", "#17365D", "Company Context", ["Evidence-backed answers", "Citations and uncertainty", "Employee working context"])
    rounded_box(draw, (545, 140, 955, 315), "#DCE6F1", "#17365D", "Meeting Actions", ["Transcript screening", "Commitments and decisions", "Reviewable handoffs"])
    rounded_box(draw, (1035, 140, 1445, 315), "#DCE6F1", "#17365D", "Recruiting", ["Criteria and profile scoring", "Founder feedback", "Outreach drafts"])

    rounded_box(draw, (260, 405, 1240, 600), "#F3F7FB", "#17365D", "Fastify TypeScript Application", ["Sessions and identity   REST and SSE   typed tool orchestration", "Citation validation   deterministic policy   exact-payload approval", "Cross-workflow handoffs   trace and persistence"])

    rounded_box(draw, (55, 705, 400, 870), "#F6F7F8", "#777777", "PostgreSQL pgvector", ["Company Evidence", "Conversations", "Meeting state and logs"], title_size=27, body_size=21)
    rounded_box(draw, (470, 705, 805, 870), "#F6F7F8", "#777777", "Letta App Server", ["Employee Personal Memory", "Hiring intent history", "User-scoped files"], title_size=27, body_size=21)
    rounded_box(draw, (875, 705, 1105, 870), "#F6F7F8", "#777777", "Model Providers", ["SoCLaaS Qwen", "Optional Claude", "Titan embeddings"], title_size=27, body_size=21)
    rounded_box(draw, (1175, 705, 1445, 870), "#F6F7F8", "#777777", "User Tools", ["Gmail or Outlook", "Calendar and chat", "GitHub Docs Sheets"], title_size=27, body_size=21)

    for x in (260, 750, 1240):
        arrow(draw, (x, 315), (x, 405))
    for x in (228, 638, 990, 1310):
        arrow(draw, (750, 600), (x, 705))

    image.save(DIAGRAM, quality=95)


def add_code_reference(doc, label, path):
    p = doc.add_paragraph()
    p.style = doc.styles["Body Text"]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=9.5, bold=True, color=NAVY)
    r = p.add_run(path)
    set_run_font(r, name="Courier New", size=8.8, color=GRAY)
    return p


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        r = p.add_run("Stellar Ark AI | SME Operating Agent Technical Documentation   |   ")
        set_run_font(r, size=8, color=GRAY)
        fld_char1 = OxmlElement("w:fldChar")
        fld_char1.set(qn("w:fldCharType"), "begin")
        instr_text = OxmlElement("w:instrText")
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = " PAGE "
        fld_char2 = OxmlElement("w:fldChar")
        fld_char2.set(qn("w:fldCharType"), "end")
        field_run = p.add_run()
        field_run._r.append(fld_char1)
        field_run._r.append(instr_text)
        field_run._r.append(fld_char2)
        set_run_font(field_run, size=8, color=GRAY)


def build():
    build_architecture_diagram()
    doc = setup_document()

    # Cover
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(58)
    p = doc.add_paragraph(style="Title")
    r = p.add_run("SME Operating Agent Technical Documentation")
    set_run_font(r, name="Aptos Display", size=30, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Architecture Reasoning Memory Safety and Evaluation")
    set_run_font(r, name="Aptos Display", size=17, color=GRAY)

    cover_table = doc.add_table(rows=7, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cover_table.autofit = False
    labels = ["Company", "Team code", "Hackathon", "Repository", "Document status", "Snapshot commit", "Date"]
    values = [
        "Stellar Ark AI",
        "<missing>",
        "Show Me Your Agents",
        "github.com/joshualumzy/personal_agent_long_context",
        "Evidence-grounded draft",
        "<missing>",
        "27 September 2026",
    ]
    for idx, (label, value) in enumerate(zip(labels, values)):
        a, b = cover_table.rows[idx].cells
        a.width = Inches(1.35)
        b.width = Inches(5.15)
        for cell in (a, b):
            set_cell_margins(cell, top=65, start=0, bottom=65, end=80)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        a.text = ""
        b.text = ""
        ar = a.paragraphs[0].add_run(label)
        set_run_font(ar, size=9.5, bold=True, color=GRAY)
        br = b.paragraphs[0].add_run(value)
        set_run_font(br, size=10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Document purpose")
    set_run_font(r, size=10, bold=True, color=NAVY)
    add_body(
        doc,
        "Explain how the implemented prototype retrieves company evidence, carries user-scoped memory, converts meeting transcripts into controlled follow-up, and supports founder-led recruiting while keeping consequential actions under human control.",
        after=0,
    )

    # Page 2
    page_break(doc)
    heading(doc, "Technical Summary")
    add_body(
        doc,
        "The SME Operating Agent is a TypeScript and Fastify modular monolith with three browser surfaces: Company Context, Meeting Actions, and Recruiting. The workflows share identity, model access, company evidence, and human-control patterns, but keep Company Evidence, Personal Memory, meeting state, and recruiting state in distinct stores and contracts.",
    )
    add_body(
        doc,
        "The primary runtime pattern is controlled model reasoning around typed tools. Models may select retrieval or workflow tools; deterministic server code validates inputs, screens prohibited content, checks citations, assigns action risk tiers, binds approval to an exact payload, and controls which effects can leave the application. The system does not give a model arbitrary database access or an unrestricted workplace-write tool.",
    )

    heading(doc, "Personas and Operating Context", 2)
    add_table(
        doc,
        ["Persona", "Technical role in the prototype", "Boundary"],
        [
            ["Stellar Ark AI employee", "Asks company questions and reviews meeting follow-up", "Target user; private company records are not included in the hackathon data"],
            ["Founder or hiring manager", "Confirms role criteria, reviews candidates and completes outreach", "Human judgment remains authoritative for hiring and sending"],
            ["OrgForge employee persona", "Provides a privacy-safe session identity and employee-visible demo corpus", "Synthetic demonstration and evaluation data only"],
            ["Administrator", "Configures credentials, providers, database and optional integrations", "Production SSO and fine-grained authorization are not claimed"],
        ],
        widths=[1.45, 2.9, 2.15],
        font_size=8.9,
    )
    add_body(
        doc,
        "The current roster contains 51 synthetic OrgForge employees. Jax, Priya, Chloe, Marcus, and Deepa are pinned as convenient demonstration identities. A signed session is the source of employee identity; a request cannot silently substitute a different employee identifier.",
    )

    heading(doc, "Implemented Product Surfaces", 2)
    add_table(
        doc,
        ["Surface", "Route", "Purpose"],
        [
            ["Company Context", "/", "Evidence-backed company questions, source inspection, conversation history and Personal Memory"],
            ["Meeting Actions", "/meetings", "Transcript intake, action extraction, decision conflicts, approval and handoff"],
            ["Recruiting", "/recruiting", "Role criteria, candidate pool, founder feedback, outreach and follow-up"],
        ],
        widths=[1.5, 1.15, 3.85],
        font_size=9.2,
    )

    # Page 3
    page_break(doc)
    heading(doc, "Architecture Design")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(DIAGRAM), width=Inches(6.7))
    add_body(
        doc,
        "The application is intentionally a modular monolith for the hackathon. A single Fastify process owns HTTP validation, sessions, orchestration, model calls, tool dispatch, and Server-Sent Events. Domain modules isolate the three workflows without adding distributed-system overhead that the prototype does not need.",
    )

    heading(doc, "Component Responsibilities", 2)
    add_table(
        doc,
        ["Layer", "Responsibility"],
        [
            ["Browser", "Plain HTML, CSS and JavaScript experiences for chat, meeting work and recruiting; the main chat can embed a recruiting panel"],
            ["Fastify application", "Authentication, validation, REST and SSE, conversation context, model selection, errors and persistence"],
            ["Domain services", "S1 reasoning loop, S2 meeting state machine and policy, and S3 recruiting state machine"],
            ["Stores", "PostgreSQL and pgvector for evidence and operational state; Letta for user-scoped Personal Memory"],
            ["Providers", "SoCLaaS Qwen by default, an optional Claude-compatible gateway, Bedrock Titan embeddings, and optional workflow APIs"],
        ],
        widths=[1.55, 4.95],
        font_size=9.0,
    )

    # Page 4
    page_break(doc)
    heading(doc, "Frameworks Models and Tools")
    add_table(
        doc,
        ["Category", "Technology", "Use in this system"],
        [
            ["Application", "Node.js 22.19 or newer TypeScript Fastify", "Modular server, validated endpoints, sessions and static browser delivery"],
            ["Reasoning models", "NUS SoCLaaS Qwen qwen3.8:27b by default", "Company retrieval, meeting extraction and drafting, and recruiting judgments"],
            ["Optional model", "Claude through the configured OpenAI-compatible gateway", "Selectable Company Context model when gateway credentials are present"],
            ["Embeddings", "Amazon Titan Text Embeddings V2", "1,024-dimensional Company Evidence vectors stored in pgvector"],
            ["Evidence store", "PostgreSQL full-text search and pgvector", "Keyword and semantic retrieval, evidence chunks, links, conversations and meeting state"],
            ["Long-term memory", "Self-hosted Letta App Server", "Employee-scoped working context and founder hiring-intent history"],
            ["Optional recruiting services", "Exa Hunter Prospeo Gmail and LinkedIn inbox reader", "Public profile search, work-email lookup, read-only reply input and manual user handoff"],
            ["Optional meeting helpers", "Google free-busy and Jev", "Availability checks and a fast reader for times or decision conflicts with model fallback"],
            ["Quality tooling", "Node test runner JSDOM Playwright TypeScript and offline evaluators", "Domain, browser, holdout, regression and live-provider checks"],
        ],
        widths=[1.25, 2.05, 3.2],
        font_size=8.5,
    )

    heading(doc, "Model and Agent Design", 2)
    add_bullets(doc, [
        "S1 is a bounded tool-using loop. It normally allows up to four reasoning steps, or eight when workflow skills are registered, and forces a final answer on the last step.",
        "The first S1 step requires a tool call. Available company tools are typed search_company_knowledge and get_related_sources; recruiting tools remain unavailable until the recruiting skill is loaded.",
        "S2 uses a model to extract candidates and draft payloads, but deterministic code validates transcript quotes, risk tiers, approvals and executable effects.",
        "S3 represents hiring as explicit role, criterion, candidate, proposal and message state rather than relying on free-form chat history alone.",
        "No multi-agent orchestration is claimed. The architecture is one application with specialised domain services and a lazily loaded recruiting skill.",
    ])

    # Page 5
    page_break(doc)
    heading(doc, "Data State and Memory Design")
    heading(doc, "Short-Term Working State", 2)
    add_table(
        doc,
        ["State", "Implementation", "Purpose"],
        [
            ["Current request", "Validated request body and authenticated employee session", "Defines the active user, question and selected model"],
            ["Conversation window", "Latest six stored user or assistant turns, each clipped to 4,000 characters", "Supports follow-up questions without loading an unbounded transcript"],
            ["S1 run state", "In-memory retrieved-source map, tool-call list and run identifier", "Ensures citations can name only sources retrieved in that run"],
            ["S2 processing state", "Queued transcript batches and persisted MeetingState", "Prevents per-line model lag and tracks actions, decisions and trace events"],
            ["S3 role state", "One JSON document per role behind a RoleRepository interface", "Keeps criteria, candidates, feedback, proposals, messages and simulated time together"],
        ],
        widths=[1.45, 2.75, 2.3],
        font_size=8.8,
    )

    heading(doc, "Long-Term Memory", 2)
    add_body(
        doc,
        "Letta stores Personal Memory, not the Company Evidence corpus. Each employee identifier maps to a separate Letta agent tag. Memory file access is confined to the resolved Memory directory and a restricted tool allow-list. Inspection excludes persona, skill and Git metadata files. Background memory updates are serialized per employee to avoid race conditions.",
    )
    add_body(
        doc,
        "Personal Memory keeps concise user context with source and recorded time. Corrections can supersede earlier entries, cancellations remain visible as history, and unresolved contradictions are retained as conflicts. This context may help interpret a question, but it is displayed separately and cannot serve as a Company Evidence citation.",
    )

    heading(doc, "Storage Boundaries", 2)
    add_table(
        doc,
        ["Store", "Contains", "Must not contain"],
        [
            ["PostgreSQL", "Company Artifacts, chunks, embeddings, explicit links, conversations, employees, meetings and action logs", "Evaluation Oracle records excluded by ingestion policy"],
            ["Letta", "Employee Personal Memory and recruiting intent events", "OrgForge Company Evidence corpus"],
            ["Recruiting role files", "Per-role operational state and public professional profile fields", "Guessed email addresses or unrelated private profile data"],
            ["Offline evaluation", "Expected answers, labels and reserved Oracle data", "Runtime tool access"],
        ],
        widths=[1.35, 3.0, 2.15],
        font_size=8.9,
    )

    # Page 6
    page_break(doc)
    heading(doc, "Company Context Reasoning Loop")
    add_numbered(doc, [
        "Authenticate the employee and reject any explicit user identifier that conflicts with the signed session.",
        "Load the latest bounded conversation window and the employee's separately labelled Personal Memory when available.",
        "Require a typed tool call to search Company Evidence. The model may run focused keyword or semantic searches and may traverse explicit links only from sources already retrieved.",
        "Accumulate retrieved evidence in a server-side map and synthesise a concise response from that evidence and the session profile.",
        "Validate every citation identifier against the map. If repair cannot produce a supported answer, return an insufficient-evidence outcome instead of fabricating a source.",
        "Persist the turn with model, provider, duration, run identifier, sources, tool calls, memory status and outcome; stream progress and final validated content over SSE when requested.",
    ])

    heading(doc, "Retrieval Design", 2)
    add_table(
        doc,
        ["Mechanism", "Implementation", "Control"],
        [
            ["Keyword retrieval", "PostgreSQL tsvector full-text search over chunks", "Employee-visible admitted artifact types only"],
            ["Semantic retrieval", "Cosine similarity over Titan V2 vectors in pgvector", "Embedding model recorded with each vector"],
            ["Graph traversal", "document_links between explicit OrgForge artifact references", "Seeds must already be present in the current retrieved map"],
            ["Citation delivery", "Inline source identifiers and inspectable source metadata", "Unknown identifiers are removed or repaired before delivery"],
        ],
        widths=[1.4, 3.0, 2.1],
        font_size=9.0,
    )
    add_body(
        doc,
        "The runtime corpus currently contains 4,966 synthetic employee-visible OrgForge artifacts across Slack, email, Confluence, Jira, Zoom transcripts, pull requests, Salesforce and Zendesk. OrgForge is privacy-safe demonstration and evaluation data, not Stellar Ark AI business data or a product dependency.",
    )

    # Page 7
    page_break(doc)
    heading(doc, "Meeting Actions Reasoning and Human Control")
    add_numbered(doc, [
        "Accept transcript text typed, pasted, replayed from a scripted scenario, or loaded from an admitted Zoom transcript. Speech-to-text is outside the prototype.",
        "Screen each segment before model use. Prohibited data is withheld from storage; documented prompt-injection shapes become blocked actions and never reach extraction.",
        "Extract commitments, company questions and decisions from new segments only. A candidate is discarded if its quoted trigger does not occur in the transcript.",
        "Answer read-only company questions through S1, compare decisions with earlier meeting decisions, and retrieve evidence or missing contact and calendar fields for drafts.",
        "Apply deterministic policy, deduplicate repeated commitments, persist the action and append trace and status events.",
        "For approval-tier work, bind approval to the payload hash. Any edit increments the version, changes the hash and returns the action to proposed state.",
        "Execute only a user-controlled handoff, an explicitly simulated record, or a draft recruiting role. Escalated and blocked actions are rejected again by the executor as defence in depth.",
    ])

    heading(doc, "Risk-Calibrated Autonomy", 2)
    add_table(
        doc,
        ["Tier", "Examples", "Runtime behavior"],
        [
            ["Automatic", "Evidence-backed answer and decision conflict", "Read-only result is produced immediately and traced"],
            ["Approval", "Email message calendar ticket document spreadsheet or hiring draft", "Employee approves the exact unchanged payload before handoff"],
            ["Escalate", "Money price refund discount or contract commitment", "Named higher authority required; employee cannot execute"],
            ["Blocked", "Secret or instruction attack", "No model processing and no executor path"],
        ],
        widths=[1.05, 2.6, 2.85],
        font_size=9.0,
    )
    add_body(
        doc,
        "Everyday effects open as prefilled drafts in the employee's own Gmail or Outlook, calendar, WhatsApp or Teams, GitHub, document, or spreadsheet surface. The external tool's own signed-in user completes the send or save, so the prototype does not hold broad workplace-write credentials.",
    )

    # Page 8
    page_break(doc)
    heading(doc, "Recruiting Reasoning and Human Control")
    add_numbered(doc, [
        "Turn a typed need, uploaded job description, or known profile link into three to six must or nice-to-have criteria; the founder reviews and confirms them before search.",
        "Use Exa public-profile search when configured or clearly labelled fictional sample profiles otherwise. Judge each criterion as yes, no or unclear with a short profile-grounded reason.",
        "Assign a deterministic candidate tier from the criterion verdicts and present the pool in the Recruiting workspace or an embedded main-chat panel.",
        "Treat founder feedback as explicit state. Repeated pass reasons may create a criterion proposal, and a stalled search may create a widening proposal; neither applies until the founder accepts it.",
        "For a chosen candidate, try Hunter and then Prospeo for a work email. The system does not guess an address. Draft outreach in the founder's voice and warn if it repeats a private pass reason.",
        "Read replies through configured Gmail, a read-only LinkedIn inbox sync, or pasted text; draft scheduling or follow-up and retain the founder's explicit hiring rationale in scoped memory.",
        "Erase closed candidate records after the configured 30-day retention period in the current prototype state model.",
    ])

    heading(doc, "Recruiting Control Points", 2)
    add_table(
        doc,
        ["Decision", "Agent contribution", "Required human action"],
        [
            ["Role criteria", "Proposes and explains criteria", "Founder edits and confirms"],
            ["Candidate assessment", "Scores declared criteria from public profile fields", "Founder reviews keeps or passes"],
            ["Preference learning", "Proposes a criterion from repeated feedback", "Founder accepts or declines"],
            ["Search expansion", "Proposes the next widening step after inactivity", "Founder accepts or declines"],
            ["Outreach", "Finds a non-guessed work email and drafts text", "Founder resolves warnings and sends from their own account"],
            ["Hiring outcome", "Tracks status and rationale", "Founder makes the decision"],
        ],
        widths=[1.55, 3.15, 1.8],
        font_size=8.8,
    )

    # Page 9
    page_break(doc)
    heading(doc, "Guardrails Security and Constraints")
    add_table(
        doc,
        ["Control area", "Implemented control", "Current claim boundary"],
        [
            ["Identity", "HMAC-SHA256 session tokens, HttpOnly SameSite cookies and scrypt password verification", "Prototype identity; enterprise SSO and fine-grained department authorization are deferred"],
            ["Input safety", "Request-size validation, prohibited-data detection and S2 pre-model injection screening", "No claim of complete prompt-injection resistance"],
            ["Evidence integrity", "Runtime source admission rules, Oracle exclusion and current-run citation validation", "A valid citation identifier does not alone prove perfect semantic support"],
            ["Memory isolation", "Per-employee Letta agent tags, confined memory tools and separate display", "Production correction deletion and tenant-governance controls are deferred"],
            ["Action safety", "Deterministic S2 tiers, payload-hash approval, edit invalidation and executor defence in depth", "No unattended consequential action or shared multi-party approval"],
            ["Recruiting privacy", "Public professional fields only, no guessed email, private-reason warnings and retention cleanup", "Legal basis fairness assessment and production privacy review remain deployment responsibilities"],
            ["Browser security", "Content Security Policy, no-referrer, nosniff and frame denial headers", "External penetration testing is <missing>"],
            ["Secrets", "Server-side environment configuration; credentials are not sent to the browser or tracked", "Deployment secret manager and rotation evidence are <missing>"],
        ],
        widths=[1.25, 3.25, 2.0],
        font_size=8.4,
    )

    heading(doc, "Prototype Constraints", 2)
    add_bullets(doc, [
        "The hackathon demo does not use private Stellar Ark AI records. Quantitative business impact still requires a controlled internal pilot.",
        "Meeting input is text. Audio capture and speech-to-text are outside the implemented surface.",
        "External writes are user-controlled handoffs or clearly marked simulations; existing third-party records are not edited through provider APIs.",
        "Optional services change the available path. Missing Exa uses labelled sample candidates; missing email lookup leaves outreach without an address; missing Google access keeps reply and availability features manual.",
        "Deployment URL, production hosting evidence, final provider configuration and submission snapshot commit are <missing>.",
    ])

    # Page 10
    page_break(doc)
    heading(doc, "Observability and Evaluation")
    heading(doc, "Operational Observability", 2)
    add_bullets(doc, [
        "S1 persists model, provider, duration, run identifier, retrieved sources, tool calls, memory status, outcome and conversation persistence status with each assistant turn.",
        "S2 emits live segment, action, trace and busy events and persists the full meeting state. meeting_action_log is an append-only record of action kind, tier, status, payload hash, time and structured detail.",
        "Application failure records carry operation, correlation identifier, user identifier and reason; transcript, question and Memory content are intentionally excluded from that failure shape.",
        "Evaluation Oracle data and expected answers remain in offline runners that the runtime agent cannot call.",
    ])

    heading(doc, "Saved Evaluation Evidence", 2)
    add_table(
        doc,
        ["Area", "Saved evidence", "Interpretation"],
        [
            ["Company Context", "76-question run using qwen3.8:27b: 100 percent citation integrity, 29 percent expected-evidence recall, 12 percent factual accuracy, 39.68 seconds mean latency", "Supports the citation-identifier guard in that sample; answer quality and retrieval recall remain improvement targets"],
            ["Meeting Actions", "One run dated 25 September 2026: 23 of 23 tiers correct, 9 of 9 unsafe lines blocked, 0 of 44 ordinary lines blocked, 0 actions in 29 negative cases, 1 of 1 conflict found", "Strong bounded guardrail evidence; action-kind recall was complete except email at 2 of 3"],
            ["Personal Memory", "Pilot over 20 selected mutation-heavy weekly Memora questions across two personas", "Development gate only; not a full benchmark or production reliability claim"],
            ["Recruiting", "Deterministic, holdout, regression, chaos and browser test assets exist", "Final externally reportable submission snapshot result is <missing>"],
        ],
        widths=[1.25, 3.45, 1.8],
        font_size=8.35,
    )

    heading(doc, "Submission Verification Commands", 2)
    add_table(
        doc,
        ["Command", "Purpose"],
        [
            ["npm run typecheck", "TypeScript contract check"],
            ["npm test", "Unit and regression suite"],
            ["npm run test:browser", "Browser rendering and interaction check"],
            ["npm run test:orgforge", "Runtime corpus and Oracle-isolation check"],
            ["npm run test:holdout and npm run test:lock", "Recruiting holdout and regression-lock checks"],
            ["npm run eval:orgforge, npm run eval:meetings, npm run eval:recruiting", "Live model evaluation paths"],
        ],
        widths=[2.45, 4.05],
        font_size=9.0,
    )
    add_body(
        doc,
        "A fresh clean verification run against the final submission snapshot is pending. The current checkout's typecheck cannot resolve several declared packages because the local dependency installation is incomplete; this draft therefore does not claim a new all-green run.",
        italic=True,
    )

    # Page 11
    page_break(doc)
    heading(doc, "Deployment and Operations")
    heading(doc, "Current Runtime Topology", 2)
    add_table(
        doc,
        ["Runtime", "Required configuration", "Purpose"],
        [
            ["Fastify process", "Node.js 22.19 or newer, session secret, database URL and model credentials", "Serves all three browser surfaces and APIs"],
            ["PostgreSQL with pgvector", "Docker Compose or equivalent managed database", "Evidence, embeddings, conversations, employee roster, meetings and logs"],
            ["Letta App Server", "Local or protected service endpoint and optional token", "Personal Memory and hiring-intent persistence"],
            ["Model endpoint", "SoCLaaS key or configured compatible gateway", "Reasoning and structured model calls"],
            ["AWS Bedrock", "Region ap-southeast-2 and authenticated profile", "Titan Text Embeddings V2 backfill and optional model path"],
            ["Optional workflow services", "Service-specific keys or Google OAuth", "Recruiting search, contact lookup, replies and calendar availability"],
        ],
        widths=[1.5, 3.1, 1.9],
        font_size=8.8,
    )

    heading(doc, "Local Start Sequence", 2)
    add_numbered(doc, [
        "Install declared Node and Python dependencies and create environment-specific secrets.",
        "Start PostgreSQL, apply migrations, ingest the declared synthetic corpus, and backfill embeddings when semantic retrieval is required.",
        "Start the Letta App Server and the Fastify application.",
        "Open the root company chat, the meetings route, or the recruiting route and confirm health, identity and configured provider status.",
        "Run deterministic verification and the selected live evaluation before recording deployment evidence or the demo video.",
    ])

    heading(doc, "Deployment Evidence Status", 2)
    add_table(
        doc,
        ["Evidence item", "Status"],
        [
            ["Public or judge-accessible deployment URL", "<missing>"],
            ["Hosting architecture and region", "<missing>"],
            ["Health-check capture", "<missing>"],
            ["Database migration and corpus counts", "Local commands defined; final deployed evidence <missing>"],
            ["Configured model and optional integrations used in the demo", "<missing>"],
            ["Submission snapshot commit", "<missing>"],
        ],
        widths=[3.4, 3.1],
        font_size=9.0,
    )

    # Page 12
    page_break(doc)
    heading(doc, "Implementation Traceability")
    add_table(
        doc,
        ["Concern", "Primary implementation", "Verification or evidence"],
        [
            ["HTTP sessions SSE and conversations", "src/http-app.ts src/auth.ts src/adapters/postgres-conversations.ts", "test/auth.test.ts test/conversations.test.ts test/sme-browser.test.ts"],
            ["S1 reasoning and citations", "src/soclaas-company-agent.ts src/adapters/postgres-company-knowledge.ts", "test/soclaas-company-agent.test.ts test/unified-agent.test.ts eval/orgforge"],
            ["Personal Memory", "src/adapters/letta-memory.ts src/memory-queue.ts", "test/letta-memory.test.ts test/memory-updates.test.ts eval/memora"],
            ["Meeting extraction and policy", "src/meetings/extractor.ts guard.ts policy.ts service.ts", "test/meetings-core.test.ts test/meetings-routes.test.ts eval/meetings"],
            ["Approval and handoff", "src/meetings/service.ts executor.ts handoff.ts", "test/meetings-store-executor.test.ts test/meetings-handoff.test.ts"],
            ["Recruiting state and tools", "src/recruiting/service.ts routes.ts chat-tools.ts roles.ts", "test/recruiting.test.ts test/holdout test/regression eval/recruiting"],
            ["Schema and storage", "database/migrations scripts/orgforge", "npm run db:migrate npm run test:orgforge"],
            ["Browser surfaces", "public/index.html app.js meetings.html meetings.js recruiting.html recruiting.js", "test/sme-browser.test.ts and browser-specific regression assets"],
        ],
        widths=[1.4, 3.15, 1.95],
        font_size=8.3,
    )

    heading(doc, "Technical Conclusion", 2)
    add_body(
        doc,
        "The prototype demonstrates a coherent context-to-action architecture: evidence retrieval is inspectable, Personal Memory is separately scoped, meeting work is risk-tiered and approval-bound, and recruiting remains founder-controlled. Its strongest design choice is that model reasoning operates inside deterministic technical boundaries rather than replacing them. The remaining submission work is operational evidence: restore the declared dependency installation, run the final verification suite, record the deployed configuration, and replace the remaining <missing> fields with the exact submission snapshot.",
    )

    add_footer(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "SME Operating Agent Technical Documentation"
    doc.core_properties.subject = "Show Me Your Agents Hackathon technical documentation"
    doc.core_properties.author = "Stellar Ark AI Hackathon Team"
    doc.core_properties.keywords = "Stellar Ark AI, SME, agent, architecture, memory, human in the loop, hackathon"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
